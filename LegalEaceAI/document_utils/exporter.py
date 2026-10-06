from io import BytesIO

from docx import Document
from fpdf import FPDF


def make_txt(text: str) -> bytes:
    """Create a UTF-8 TXT file."""
    return text.encode("utf-8")


def make_docx(text: str, document_type: str = "Legal Document") -> bytes:
    """Create a DOCX document from the generated legal text."""
    document = Document()

    document.add_heading(document_type, level=1)

    for paragraph in text.split("\n\n"):
        paragraph = paragraph.strip()

        if paragraph:
            document.add_paragraph(paragraph)

    output = BytesIO()
    document.save(output)

    return output.getvalue()


def make_pdf(text: str, document_type: str = "Legal Document") -> bytes:
    """Create a PDF document from the generated legal text."""
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, document_type)

    pdf.ln(5)

    pdf.set_font("Helvetica", size=11)

    for paragraph in text.split("\n\n"):
        paragraph = paragraph.strip()

        if paragraph:
            pdf.multi_cell(0, 7, paragraph)
            pdf.ln(3)

    return bytes(pdf.output())