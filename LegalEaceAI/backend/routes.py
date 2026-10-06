from io import BytesIO

from fastapi import APIRouter, HTTPException
from fastapi.responses import Response

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from fpdf import FPDF

from ai_core.gemini_generator import GeminiDocumentGenerator
from backend.models import (
    GenerateRequest,
    GenerateResponse,
    DownloadRequest,
)


router = APIRouter()


# =========================================================
# GENERATE DOCUMENT
# =========================================================

@router.post(
    "/generate",
    response_model=GenerateResponse
)
def generate_document(
    request: GenerateRequest
) -> GenerateResponse:

    try:
        generator = GeminiDocumentGenerator()

        content = generator.generate_document(
            document_type=request.document_type,
            parties=request.parties,
            terms=request.terms,
            effective_date=request.effective_date,
            jurisdiction=request.jurisdiction,
            language=request.language,
        )

        return GenerateResponse(
            content=content,
            document_type=request.document_type,
            model=generator.model,
        )

    except RuntimeError as exc:
        raise HTTPException(
            status_code=500,
            detail=str(exc)
        ) from exc

    except Exception as exc:
        raise HTTPException(
            status_code=502,
            detail=f"AI generation failed: {exc}"
        ) from exc


# =========================================================
# DOWNLOAD TXT
# =========================================================

@router.post("/download/txt")
def download_txt(request: DownloadRequest):

    filename = (
        request.document_type
        .replace(" ", "_")
        .replace("/", "_")
        + ".txt"
    )

    return Response(
        content=request.content,
        media_type="text/plain; charset=utf-8",
        headers={
            "Content-Disposition":
                f'attachment; filename="{filename}"'
        }
    )


# =========================================================
# DOWNLOAD DOCX
# =========================================================

@router.post("/download/docx")
def download_docx(request: DownloadRequest):

    try:
        document = Document()

        # Page margins
        section = document.sections[0]
        section.top_margin = Pt(60)
        section.bottom_margin = Pt(60)
        section.left_margin = Pt(65)
        section.right_margin = Pt(65)

        # Default font
        styles = document.styles
        styles["Normal"].font.name = "Times New Roman"
        styles["Normal"].font.size = Pt(12)

        # Title
        title = document.add_paragraph()

        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = title.add_run(
            request.document_type.upper()
        )

        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(16)

        # Document content
        paragraphs = request.content.split("\n")

        for line in paragraphs:

            line = line.strip()

            if not line:
                document.add_paragraph()
                continue

            paragraph = document.add_paragraph()

            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.line_spacing = 1.5

            run = paragraph.add_run(line)

            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

            # Heading detection
            if (
                line.isupper()
                and len(line) < 100
            ):
                run.bold = True
                run.font.size = Pt(13)

        # Footer
        footer = section.footer

        footer_paragraph = footer.paragraphs[0]
        footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        footer_run = footer_paragraph.add_run(
            "LegalEase AI • AI-assisted legal document draft"
        )

        footer_run.font.name = "Times New Roman"
        footer_run.font.size = Pt(9)

        # Save to memory
        output = BytesIO()

        document.save(output)

        output.seek(0)

        filename = (
            request.document_type
            .replace(" ", "_")
            .replace("/", "_")
            + ".docx"
        )

        return Response(
            content=output.getvalue(),
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            headers={
                "Content-Disposition":
                    f'attachment; filename="{filename}"'
            }
        )

    except Exception as exc:

        raise HTTPException(
            status_code=500,
            detail=f"DOCX generation failed: {exc}"
        ) from exc


# =========================================================
# DOWNLOAD PDF
# =========================================================

@router.post("/download/pdf")
def download_pdf(request: DownloadRequest):

    try:

        pdf = FPDF()

        pdf.set_auto_page_break(
            auto=True,
            margin=20
        )

        pdf.add_page()

        # Title
        pdf.set_font(
            "Times",
            "B",
            16
        )

        pdf.multi_cell(
            0,
            10,
            request.document_type.upper(),
            align="C"
        )

        pdf.ln(8)

        # Body
        pdf.set_font(
            "Times",
            "",
            12
        )

        # FPDF built-in fonts support limited characters.
        # Replace unsupported characters safely.
        safe_content = request.content.encode(
            "latin-1",
            errors="replace"
        ).decode("latin-1")

        for paragraph in safe_content.split("\n"):

            paragraph = paragraph.strip()

            if not paragraph:
                pdf.ln(5)
                continue

            pdf.multi_cell(
                0,
                7,
                paragraph
            )

            pdf.ln(2)

        # Footer
        pdf.set_y(-15)

        pdf.set_font(
            "Times",
            "I",
            8
        )

        pdf.cell(
            0,
            10,
            "LegalEase AI - AI-assisted legal document draft",
            align="C"
        )

        pdf_bytes = bytes(
            pdf.output()
        )

        filename = (
            request.document_type
            .replace(" ", "_")
            .replace("/", "_")
            + ".pdf"
        )

        return Response(
            content=pdf_bytes,
            media_type="application/pdf",
            headers={
                "Content-Disposition":
                    f'attachment; filename="{filename}"'
            }
        )

    except Exception as exc:

        raise HTTPException(
            status_code=500,
            detail=f"PDF generation failed: {exc}"
        ) from exc